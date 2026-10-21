"""tests/test_efuel_economics.py — tests du module economics/efuel_economics.py"""
import math
import os

from economics.efuel_economics import (
    Params, PRODUCTS, crf, lcoh, product_cost, plant_sizing, stoichiometry,
)


def test_crf_known_value():
    # 8 %, 20 ans -> 0.10185
    assert math.isclose(crf(0.08, 20), 0.101852, rel_tol=1e-4)
    assert math.isclose(crf(0.0, 20), 1 / 20)


def test_stoichiometry_methanol():
    s = stoichiometry("e-methanol")          # CO2 + 3 H2 -> CH3OH + H2O
    assert math.isclose(s["h2"], 0.1888, rel_tol=1e-3)
    assert math.isclose(s["co2"], 1.3736, rel_tol=1e-3)


def test_stoichiometry_ammonia_has_no_co2():
    s = stoichiometry("e-ammonia")
    assert s["co2"] == 0.0 and math.isclose(s["n2"], 0.8225, rel_tol=1e-3)


def test_lcoh_in_article_range():
    # L'article cite 3-6 EUR/kg selon électrolyseur et prix de l'électricité
    assert 3.0 <= lcoh(Params())["total"] <= 6.0


def test_lcoh_decreases_with_flh_and_cheaper_power():
    base = lcoh(Params())["total"]
    assert lcoh(Params(flh=8000))["total"] < base
    assert lcoh(Params(elec_price_eur_mwh=20))["total"] < base


def test_all_products_positive_and_efficiency_plausible():
    for k in PRODUCTS:
        c = product_cost(Params(), k)
        assert c["total_eur_t"] > 0
        assert 0.30 < c["efficiency"] < 0.60     # chaîne électricité -> carburant
        assert 0.5 < c["h2_share"] < 0.95


def test_dac_costs_more_than_point_source():
    a = product_cost(Params(co2_type="point_source"), "e-kerosene")["total_eur_t"]
    b = product_cost(Params(co2_type="dac"), "e-kerosene")["total_eur_t"]
    assert b > a


def test_ammonia_independent_of_co2_type():
    a = product_cost(Params(co2_type="point_source"), "e-ammonia")["total_eur_t"]
    b = product_cost(Params(co2_type="dac"), "e-ammonia")["total_eur_t"]
    assert math.isclose(a, b)


def test_sizing_matches_benchmarks_order_of_magnitude():
    # Projets de l'article : ~2,9-3,3 MW d'électrolyse par kt/an (e-hydrocarbures, ~8000 h/an)
    s = plant_sizing(Params(flh=8000), "e-kerosene", 1000.0)
    assert 2.8 <= s["electrolyser_mw"] <= 3.6


def test_excel_crosscheck_if_available():
    """Compare au classeur recalculé (si présent). Chemin : EFUELS_XLSX."""
    path = os.environ.get("EFUELS_XLSX")
    if not path or not os.path.exists(path):
        return
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Calculateur_LCOE"]
    names = ["e-methanol", "e-DME", "e-kerosene", "e-diesel", "e-gasoline", "e-ammonia"]
    for i, n in enumerate(names):
        assert math.isclose(ws[f"F{71+i}"].value, product_cost(Params(), n)["total_eur_t"], rel_tol=1e-6)
